"""Prepare complete, source-addressable reading units; never equate parsing with reading."""

import csv
import io
import json
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any, Dict, Optional


def _decode(raw: bytes) -> str:
    if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        return raw.decode("utf-32")
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    if b"\x00" in raw:
        raise ValueError("二进制或未声明编码的文本，不作有损转录")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("gb18030")


def prepare_document(source: Path, output_dir: Path, document_id: str,
                     *, max_source_bytes: int = 100 * 1024 * 1024,
                     max_units: int = 10000, chunk_chars: int = 8000,
                     render_dpi: int = 144, max_page_pixels: int = 40_000_000,
                     max_sheet_cells: int = 2_000_000) -> Dict[str, Any]:
    """Return a manifest. ``ready`` means prepared, never that an LLM has read it.

    All limits fail explicitly. Inputs and referenced external attachments are never
    downloaded or overwritten. Unsupported/partially prepared content is incomplete.
    Each output directory should belong to one document, as enforced by the caller.
    """
    source, output_dir = Path(source).resolve(), Path(output_dir).resolve()
    if min(max_source_bytes, max_units, chunk_chars, render_dpi,
           max_page_pixels, max_sheet_cells) <= 0:
        raise ValueError("文档准备资源限制必须为正数")
    if not re.fullmatch(r"[\w.-]{1,128}", document_id, re.UNICODE) or document_id in (".", ".."):
        raise ValueError("document_id 必须是安全且非空的标识")
    result = {"document_id": document_id, "format": source.suffix.lower().lstrip("."),
              "status": "ready", "units": [], "warnings": []}

    def warn(message: str, incomplete: bool = True) -> None:
        result["warnings"].append(message)
        if incomplete:
            result["status"] = "incomplete"

    def add(unit_id: str, text: str, locator: str,
            image_path: Optional[Path] = None) -> None:
        if len(result["units"]) >= max_units:
            raise ValueError("阅读单元数超过 max_units=%d，剩余内容未准备" % max_units)
        result["units"].append({"unit_id": unit_id, "document_id": document_id,
                                "ordinal": len(result["units"]) + 1, "text": text,
                                "image_path": str(image_path) if image_path else None,
                                "locator": locator})

    def chunks(text: str, prefix: str = "t", locator: str = "字符") -> None:
        for start in range(0, len(text), chunk_chars):
            add("%s%05d" % (prefix, start // chunk_chars + 1),
                text[start:start + chunk_chars], "%s %d–%d" %
                (locator, start + 1, min(start + chunk_chars, len(text))))

    try:
        if not source.is_file():
            raise ValueError("原件不存在或不是文件：%s" % source)
        if source.stat().st_size > max_source_bytes:
            raise ValueError("原件超过 max_source_bytes=%d，未截断读取" % max_source_bytes)
        output_dir.mkdir(parents=True, exist_ok=True)
        suffix = source.suffix.lower()
        if suffix in (".txt", ".md", ".json", ".jsonl", ".xml", ".html", ".htm", ".csv", ".tsv"):
            raw_text = _decode(source.read_bytes())
            if suffix in (".csv", ".tsv"):
                # 用 JSON 数组保留空格、换行、空单元格及前导零，不推测数据类型。
                reader = csv.reader(io.StringIO(raw_text, newline=""),
                                    delimiter="\t" if suffix == ".tsv" else ",", strict=True)
                for row_number, row in enumerate(reader, 1):
                    encoded = json.dumps(row, ensure_ascii=False)
                    for start in range(0, len(encoded), chunk_chars):
                        add("r%06d-%03d" % (row_number, start // chunk_chars + 1),
                            encoded[start:start + chunk_chars],
                            "逻辑行 %d；数组字符 %d–%d（跨单元续接）" %
                            (row_number, start + 1, min(start + chunk_chars, len(encoded))))
            else:
                chunks(raw_text)
            if suffix == ".md" and re.search(r"!\[|<\s*(?:img|svg|iframe)\b", raw_text, re.I):
                warn("Markdown 含图片或嵌入资源；正文完整保留，视觉内容需另行加入材料包")
            if suffix in (".html", ".htm", ".xml"):
                if re.search(r"<\s*(?:img|svg|canvas|object|embed|iframe|video|audio|script|style)\b|"
                             r"<\s*link\b[^>]*stylesheet|"
                             r"(?:background(?:-image)?\s*:|url\s*\()", raw_text, re.I):
                    warn("网页含视觉、动态或嵌入内容；原始标记完整保留，但尚未渲染这些内容")
                if re.search(r"\bdownload(?:\s|=|>)", raw_text, re.I):
                    warn("网页含下载附件入口；附件尚未加入材料包")
        elif suffix == ".pdf":
            _prepare_pdf(source, output_dir, add, warn, render_dpi, max_page_pixels)
        elif suffix in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff"):
            from PIL import Image, ImageOps
            with Image.open(source) as original:
                for frame in range(getattr(original, "n_frames", 1)):
                    original.seek(frame)
                    if original.width * original.height > max_page_pixels:
                        raise ValueError("图像第 %d 帧超过像素限制，未缩略替代全文" % (frame + 1))
                    oriented = ImageOps.exif_transpose(original.copy()).convert("RGBA")
                    background = Image.new("RGBA", oriented.size, "white")
                    image = Image.alpha_composite(background, oriented).convert("RGB")
                    image_path = output_dir / ("%s-f%04d.png" % (document_id, frame + 1))
                    if image_path == source:
                        raise ValueError("输出路径不能覆盖原件")
                    image.save(image_path, format="PNG")
                    add("f%04d" % (frame + 1), "图像帧 %d；原始尺寸 %d×%d" %
                        (frame + 1, original.width, original.height),
                        "图像第 %d 帧" % (frame + 1), image_path)
        elif suffix == ".xlsx":
            _prepare_xlsx(source, add, warn, chunk_chars, max_sheet_cells)
        else:
            warn("暂不支持 %s 的完整准备；请提供 PDF、文本、CSV、XLSX 或逐页图片，原件保留" %
                 (suffix or "无扩展名文件"))
        if not result["units"]:
            warn("没有可阅读单元；不能认定已准备完整或业务不适用")
    except Exception as exc:
        warn("文档准备失败 [%s]：%s" % (type(exc).__name__, exc))
    return result


def _prepare_pdf(source, output_dir, add, warn, render_dpi, max_page_pixels):
    from pypdf import PdfReader
    reader = PdfReader(str(source), strict=False)
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("PDF 已加密且无法用空密码读取")
    root = reader.trailer["/Root"]
    names = root.get("/Names")
    if names and names.get_object().get("/EmbeddedFiles"):
        warn("PDF 含内嵌附件，尚未展开加入材料包")
    if root.get("/AF"):
        warn("PDF 含关联文件，尚未展开加入材料包")
    if root.get("/AcroForm"):
        warn("PDF 含交互表单，隐藏字段或尚未生成外观的值未独立展开")
    if root.get("/OCProperties"):
        warn("PDF 含可选内容图层，默认页面图像不能证明隐藏图层已经展开")
    rendered = None
    try:
        try:
            import fitz
            rendered = fitz.open(str(source))
            if rendered.needs_pass and not rendered.authenticate(""):
                raise ValueError("PDF 图像渲染需要密码")
            if len(rendered) != len(reader.pages):
                warn("PDF 文本与渲染页数不一致")
        except ImportError:
            warn("缺少 PyMuPDF，无法提供 PDF 全页图像；纯文本转录不足以认定全文准备完整")
        for index, page in enumerate(reader.pages):
            text, image_path = "", None
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                warn("PDF 第 %d 页文本转录失败：%s；仍尝试原页图像" % (index + 1, exc), False)
            try:
                for reference in page.get("/Annots", []):
                    annotation = reference.get_object()
                    if annotation.get("/Subtype") == "/FileAttachment":
                        warn("PDF 第 %d 页含文件附件，尚未展开" % (index + 1))
                    if annotation.get("/Contents"):
                        text += "\n[原页批注内容] " + str(annotation["/Contents"])
            except Exception as exc:
                warn("PDF 第 %d 页批注读取失败：%s" % (index + 1, exc))
            if rendered is not None:
                try:
                    visual_page = rendered.load_page(index)
                    zoom = render_dpi / 72.0
                    if visual_page.rect.width * visual_page.rect.height * zoom * zoom > max_page_pixels:
                        raise ValueError("页面超过像素限制，未缩略替代原页")
                    pixmap = visual_page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
                    image_path = output_dir / ("p%04d.png" % (index + 1))
                    pixmap.save(str(image_path))
                except Exception as exc:
                    image_path = None
                    warn("PDF 第 %d 页图像准备失败：%s" % (index + 1, exc))
            add("p%04d" % (index + 1), text, "PDF 物理页 %d" % (index + 1), image_path)
    finally:
        if rendered is not None:
            rendered.close()


def _prepare_xlsx(source, add, warn, chunk_chars, max_sheet_cells):
    import openpyxl
    with zipfile.ZipFile(source) as archive:
        names = archive.namelist()
        if sum(info.file_size for info in archive.infolist()) > 512 * 1024 * 1024:
            raise ValueError("XLSX 解压后超过 512 MiB，未继续加载")
        visuals = [name for name in names if name.startswith(
            ("xl/media/", "xl/charts/", "xl/embeddings/", "xl/drawings/"))]
        # Excel 评论也使用 VML 绘图；仅当所有形状均为文本批注时可由 comment 字段完整表达。
        for name in list(visuals):
            if name.endswith(".vml"):
                tree = ET.fromstring(archive.read(name))
                shapes = [node for node in tree.iter() if node.tag.split("}")[-1] == "shape"]
                image_nodes = [node for node in tree.iter() if node.tag.split("}")[-1] in ("image", "imagedata")]
                if shapes and not image_nodes and all(any(
                        child.tag.split("}")[-1] == "ClientData" and child.get("ObjectType") == "Note"
                        for child in shape.iter()) for shape in shapes):
                    visuals.remove(name)
        if visuals:
            warn("XLSX 含图像、图表、绘图或内嵌对象；单元格已准备，视觉内容需另行导出")
        if any(name.startswith("xl/externalLinks/") for name in names):
            warn("XLSX 含外部链接；外部依赖未加入材料包，缓存不代表已重新计算")
    formula_book = openpyxl.load_workbook(source, data_only=False, read_only=False)
    value_book = None
    try:
        value_book = openpyxl.load_workbook(source, data_only=True, read_only=False)
        for sheet_number, sheet in enumerate(formula_book.worksheets, 1):
            if sheet.max_row * sheet.max_column > max_sheet_cells:
                warn("工作表 %s 超过单元格扫描上限 %d，未静默缩小范围" % (sheet.title, max_sheet_cells))
                continue
            if len(sheet.conditional_formatting):
                warn("工作表 %s 含条件格式，视觉含义未渲染；请另附该表 PDF 或图像" % sheet.title)
            names = (formula_book.defined_names.values() if hasattr(formula_book.defined_names, "values")
                     else formula_book.defined_names.definedName)
            metadata = {"sheet": sheet.title, "state": sheet.sheet_state,
                        "dimensions": sheet.calculate_dimension(),
                        "merged_cells": [str(region) for region in sheet.merged_cells.ranges],
                        "hidden_rows": [number for number, dim in sheet.row_dimensions.items() if dim.hidden],
                        "hidden_columns": [name for name, dim in sheet.column_dimensions.items() if dim.hidden],
                        "defined_names": [{"name": name.name, "definition": name.attr_text,
                                           "local_sheet_id": name.localSheetId} for name in names],
                        "headers_footers": {name: str(getattr(sheet, name)) for name in
                                            ("oddHeader", "oddFooter", "evenHeader", "evenFooter",
                                             "firstHeader", "firstFooter")},
                        "formula_notice": "公式与现有缓存分别保留；缓存是否最新未经重算验证"}
            records = [("工作表元数据", metadata)]
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None and not cell.comment and not cell.hyperlink:
                        continue
                    item = {"coordinate": cell.coordinate, "value": cell.value,
                            "data_type": cell.data_type, "number_format": cell.number_format}
                    if cell.data_type == "f":
                        if not isinstance(cell.value, str):
                            warn("%s!%s 含非普通字符串公式，尚未完整解释公式对象" % (sheet.title, cell.coordinate))
                        item["formula"] = cell.value
                        item["cached_value"] = value_book[sheet.title][cell.coordinate].value
                        item["cache_status"] = ("missing" if item["cached_value"] is None
                                                else "present_unverified")
                    if cell.comment:
                        item["comment"] = cell.comment.text
                    if cell.hyperlink:
                        item["hyperlink"] = cell.hyperlink.target or cell.hyperlink.location
                    records.append((cell.coordinate, item))
            # 每个单元格有原坐标；大单元格分块保留全部内容，并显式标识续接。
            buffer, start_locator, segment = [], "", 1
            for locator, record in records:
                encoded = json.dumps(record, ensure_ascii=False, default=str) + "\n"
                if buffer and sum(len(item) for item in buffer) + len(encoded) > chunk_chars:
                    add("s%03d-%05d" % (sheet_number, segment), "".join(buffer),
                        "%s!%s 起" % (sheet.title, start_locator))
                    segment += 1
                    buffer = []
                if len(encoded) > chunk_chars:
                    for start in range(0, len(encoded), chunk_chars):
                        add("s%03d-%05d" % (sheet_number, segment), encoded[start:start + chunk_chars],
                            "%s!%s；记录字符 %d 起（续接）" % (sheet.title, locator, start + 1))
                        segment += 1
                else:
                    if not buffer:
                        start_locator = locator
                    buffer.append(encoded)
            if buffer:
                add("s%03d-%05d" % (sheet_number, segment), "".join(buffer),
                    "%s!%s 起" % (sheet.title, start_locator))
    finally:
        formula_book.close()
        if value_book is not None:
            value_book.close()
