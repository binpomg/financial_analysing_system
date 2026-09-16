import csv
import importlib.util
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from finresearch.documents import prepare_document


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def prepare_text(self, text, suffix=".txt", **kwargs):
        source = self.root / ("source" + suffix)
        with source.open("w", encoding="utf-8", newline="") as output:
            output.write(text)
        return prepare_document(source, self.root / "prepared", "test-document", **kwargs)

    def test_long_text_keeps_exact_content_and_stable_units(self):
        text = "第一行\r\n附注：单位万元。\n" * 3000
        result = self.prepare_text(text, chunk_chars=101)
        self.assertEqual(result["status"], "ready")
        self.assertEqual("".join(unit["text"] for unit in result["units"]), text)
        ids = [unit["unit_id"] for unit in result["units"]]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertEqual(result["units"][-1]["ordinal"], len(ids))

    def test_gb18030_is_decoded_without_replacement(self):
        source = self.root / "legacy.txt"
        source.write_bytes("附注：未披露，单位万元".encode("gb18030"))
        result = prepare_document(source, self.root / "prepared", "legacy")
        self.assertEqual(result["units"][0]["text"], "附注：未披露，单位万元")

    def test_limit_marks_incomplete_instead_of_silent_truncation(self):
        result = self.prepare_text("123456789", chunk_chars=2, max_units=2)
        self.assertEqual(result["status"], "incomplete")
        self.assertEqual(len(result["units"]), 2)
        self.assertTrue(any("max_units" in warning for warning in result["warnings"]))

    def test_missing_file_is_structured_incomplete(self):
        result = prepare_document(self.root / "missing.pdf", self.root / "prepared", "missing")
        self.assertEqual(result["status"], "incomplete")
        self.assertEqual(result["units"], [])

    def test_empty_file_is_not_full_readiness(self):
        self.assertEqual(self.prepare_text("")["status"], "incomplete")

    def test_csv_preserves_multiline_cells_blank_cells_and_leading_zero(self):
        stream = io.StringIO(newline="")
        rows = [["主体", "内容", "代码", "空字段"], ["公司甲", "第一行\n脚注第二行", "000001", ""]]
        writer = csv.writer(stream)
        writer.writerows(rows)
        result = self.prepare_text(stream.getvalue(), ".csv")
        self.assertEqual(result["status"], "ready")
        self.assertEqual([json.loads(unit["text"]) for unit in result["units"]], rows)

    def test_html_visual_content_is_not_silently_discarded(self):
        text = '<html><table><tr><td>金额</td></tr></table><img src="scan.png"></html>'
        result = self.prepare_text(text, ".html")
        self.assertEqual(result["status"], "incomplete")
        self.assertEqual("".join(unit["text"] for unit in result["units"]), text)

    def test_markdown_image_is_explicit_incomplete(self):
        result = self.prepare_text("![扫描表](table.png)\n附件说明", ".md")
        self.assertEqual(result["status"], "incomplete")

    def test_external_stylesheet_does_not_claim_visual_completeness(self):
        result = self.prepare_text('<link rel="stylesheet" href="table.css"><p>100</p>', ".html")
        self.assertEqual(result["status"], "incomplete")

    def test_unsupported_word_file_is_not_claimed_readable(self):
        result = self.prepare_text("placeholder", ".docx")
        self.assertEqual(result["status"], "incomplete")

    @unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow is optional")
    def test_all_tiff_frames_are_prepared(self):
        from PIL import Image
        source = self.root / "scans.tiff"
        page_one = Image.new("RGB", (50, 40), "white")
        page_two = Image.new("RGB", (60, 50), "gray")
        page_one.save(source, save_all=True, append_images=[page_two])
        result = prepare_document(source, self.root / "prepared", "scans")
        self.assertEqual(result["status"], "ready")
        self.assertEqual(len(result["units"]), 2)
        self.assertTrue(all(Path(unit["image_path"]).exists() for unit in result["units"]))

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "pypdf is optional")
    def test_pdf_without_visual_renderer_is_incomplete(self):
        from pypdf import PdfWriter
        source = self.root / "pages.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.add_blank_page(width=200, height=200)
        with source.open("wb") as output:
            writer.write(output)
        with mock.patch.dict("sys.modules", {"fitz": None}):
            result = prepare_document(source, self.root / "prepared", "pdf")
        self.assertEqual(result["status"], "incomplete")
        self.assertEqual([unit["unit_id"] for unit in result["units"]], ["p0001", "p0002"])
        self.assertTrue(any("PyMuPDF" in warning for warning in result["warnings"]))

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "pypdf is optional")
    def test_pdf_embedded_attachment_is_reported(self):
        from pypdf import PdfWriter
        source = self.root / "attachment.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.add_attachment("补充资料.txt", b"Additional source")
        with source.open("wb") as output:
            writer.write(output)
        result = prepare_document(source, self.root / "prepared", "attachment")
        self.assertEqual(result["status"], "incomplete")
        self.assertTrue(any("内嵌附件" in warning for warning in result["warnings"]))

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "pypdf is optional")
    def test_pdf_hidden_form_values_require_explicit_preparation(self):
        from pypdf import PdfWriter
        from pypdf.generic import ArrayObject, DictionaryObject, NameObject, TextStringObject
        source = self.root / "form.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        field = DictionaryObject({NameObject("/FT"): NameObject("/Tx"),
                                  NameObject("/T"): TextStringObject("hidden_value"),
                                  NameObject("/V"): TextStringObject("critical source value")})
        writer._root_object[NameObject("/AcroForm")] = DictionaryObject({
            NameObject("/Fields"): ArrayObject([field])})
        with source.open("wb") as output:
            writer.write(output)
        result = prepare_document(source, self.root / "prepared", "form")
        self.assertEqual(result["status"], "incomplete")
        self.assertTrue(any("交互表单" in warning for warning in result["warnings"]))

    @unittest.skipUnless(importlib.util.find_spec("fitz") and importlib.util.find_spec("pypdf"),
                         "PDF rendering dependencies are optional")
    def test_pdf_prepares_every_visual_page_and_footnote(self):
        import fitz
        source = self.root / "complete.pdf"
        pdf = fitz.open()
        for number in range(2):
            page = pdf.new_page(width=300, height=400)
            page.insert_text((20, 40), "Page %d revenue: 100" % (number + 1))
            page.insert_text((20, 380), "Footnote: CNY ten thousand")
        pdf.save(str(source))
        pdf.close()
        result = prepare_document(source, self.root / "prepared", "complete")
        self.assertEqual(result["status"], "ready", result["warnings"])
        self.assertEqual(len(result["units"]), 2)
        for unit in result["units"]:
            self.assertIn("Footnote", unit["text"])
            self.assertTrue(Path(unit["image_path"]).is_file())

    @unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow is optional")
    def test_transparent_image_is_rendered_on_readable_white_background(self):
        from PIL import Image
        source = self.root / "transparent.png"
        original = Image.new("RGBA", (10, 10), (0, 0, 0, 0))
        original.putpixel((5, 5), (0, 0, 0, 255))
        original.save(source)
        result = prepare_document(source, self.root / "prepared", "transparent")
        with Image.open(result["units"][0]["image_path"]) as image:
            self.assertEqual(image.getpixel((0, 0)), (255, 255, 255))
            self.assertEqual(image.getpixel((5, 5)), (0, 0, 0))

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is optional")
    def test_xlsx_preserves_formula_merged_hidden_and_comment(self):
        import openpyxl
        from openpyxl.comments import Comment
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "财务"
        sheet["A1"] = "单位：万元"
        sheet.merge_cells("A1:C1")
        sheet["A2"] = "营业收入"
        sheet["B2"] = 120
        sheet["B2"].comment = Comment("含税口径", "author")
        sheet["C2"] = "=B2*2"
        sheet.row_dimensions[2].hidden = True
        sheet.column_dimensions["C"].hidden = True
        hidden = book.create_sheet("隐藏附注")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "仍需读取"
        source = self.root / "financial.xlsx"
        book.save(source)
        book.close()
        result = prepare_document(source, self.root / "prepared", "financial")
        self.assertEqual(result["status"], "ready", result["warnings"])
        text = "".join(unit["text"] for unit in result["units"])
        self.assertIn('"merged_cells": ["A1:C1"]', text)
        self.assertIn('"hidden_rows": [2]', text)
        self.assertIn('"hidden_columns": ["C"]', text)
        self.assertIn('"formula": "=B2*2"', text)
        self.assertIn('"cache_status": "missing"', text)
        self.assertIn("含税口径", text)
        self.assertIn("仍需读取", text)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is optional")
    def test_xlsx_chart_requires_visual_preparation(self):
        import openpyxl
        from openpyxl.chart import BarChart, Reference
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(["季度", "收入"])
        sheet.append(["Q1", 100])
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        sheet.add_chart(chart, "D1")
        source = self.root / "chart.xlsx"
        book.save(source)
        book.close()
        result = prepare_document(source, self.root / "prepared", "chart")
        self.assertEqual(result["status"], "incomplete")
        self.assertTrue(any("图表" in warning for warning in result["warnings"]))


if __name__ == "__main__":
    unittest.main()
